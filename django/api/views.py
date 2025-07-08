import base64

from django.db import transaction
from django.db.models import CharField, Count, Q, Value
from django.db.models.functions import Concat
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from djoser.views import TokenCreateView, UserViewSet
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from rest_framework import filters, generics, status, viewsets
from rest_framework.authtoken.models import Token
from rest_framework.decorators import action
from rest_framework.mixins import ListModelMixin, RetrieveModelMixin
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from employees.models import Competence, Employee, Organization, Rating
from homepage.models import Answer, News, Poll, PollGroup, PollSubmission, Question

from .filters import CompetenceFilter
from .permissions import IsAdminUserOrReadOnly, IsUserOrReadOnly
from .serializers import (
    CompetenceSerializer,
    FileUploadSerializer,
    HierarchySerializer,
    MyProfileSerializer,
    NewsSerializer,
    OrganizationSerializer,
    OrgStructureSerializer,
    PollGroupSerializer,
    PollSerializer,
    PollSubmissionCreateSerializer,
    PollSubmissionWithAnswersSerializer,
    ProfileInOrganizationSerializer,
    RatingDELETESerializer,
    RatingPOSTSerializer,
    RatingPUTSerializer,
)


class FileUploadAPIView(APIView):
    parser_classes = (MultiPartParser, FormParser)
    serializer_class = FileUploadSerializer

    permission_classes = (IsAuthenticated,)

    def post(self, request, *args, **kwargs):
        serializer = self.serializer_class(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class PollViewset(viewsets.ModelViewSet):
    """
    Вьюсет для управления опросами и их прохождения.
    Права доступа настраиваются отдельно.
    """

    serializer_class = PollSerializer
    permission_classes = [
        IsAuthenticated,
        IsAdminUserOrReadOnly,
    ]

    filter_backends = (
        DjangoFilterBackend,
        filters.SearchFilter,
        filters.OrderingFilter,
    )
    filterset_fields = {
        "status": ["exact"],
        "kind": ["exact"],
        "poll_group__id": ["exact"],
        "organization__id": ["exact"],
        "is_public": ["exact"],
        "is_anonymous": ["exact"],
        "author__username": ["exact", "icontains"],
    }
    search_fields = ["name", "description", "questions__text"]
    ordering_fields = ["name", "created_at", "pub_date", "status"]
    ordering = ["-created_at"]

    def get_queryset(self):
        user = self.request.user
        now = timezone.now()

        if user.is_staff or user.is_superuser:
            return (
                Poll.objects.all()
                .select_related("author", "poll_group")
                .prefetch_related(
                    "questions__choices",
                    "questions__dependency_rule",
                    "organization",
                    "editors",
                    "stats_viewers",
                )
            )

        published_polls = (
            Poll.objects.filter(status=Poll.StatusChoices.PUBLISHED, pub_date__lte=now)
            .exclude(completion_date__isnull=False, completion_date__lte=now)
            .select_related("author", "poll_group")
            .prefetch_related("organization", "questions")
        )

        if user.is_authenticated:
            organization_q = Q()
            if hasattr(user, "organization") and user.organization:
                organization_q = Q(organization=user.organization)

            accessible_polls = published_polls.filter(
                Q(is_public=True) | organization_q
            ).distinct()
            return accessible_polls
        else:
            return published_polls.filter(is_public=True).distinct()

    def perform_create(self, serializer):
        serializer.save(author=self.request.user)

    def _get_answer_text(self, answer, question):
        """Helper to format an answer into a single string."""
        if not answer:
            return ""

        if question.question_type in [Question.QuestionType.SINGLE_CHOICE, Question.QuestionType.MULTIPLE_CHOICE]:
            choice_texts = [c.choice_text for c in answer.selected_choices.all()]
            if question.allow_custom_answer and answer.custom_choice_text:
                choice_texts.append(f"Свой вариант: {answer.custom_choice_text}")
            return ", ".join(choice_texts)

        return answer.free_text_answer or ""

    @action(detail=True, methods=["post"])
    def complete(self, request, pk=None):
        """Завершает опрос."""
        poll = self.get_object()
        if poll.status == Poll.StatusChoices.COMPLETED:
            return Response(
                {"message": "Опрос уже завершен."}, status=status.HTTP_400_BAD_REQUEST
            )

        if not (
            request.user.is_staff
            or poll.author == request.user
            or poll.editors.filter(id=request.user.id).exists()
        ):
            return Response(
                {"detail": "У вас нет прав для завершения этого опроса."},
                status=status.HTTP_403_FORBIDDEN,
            )

        poll.status = Poll.StatusChoices.COMPLETED
        poll.completion_date = timezone.now()
        poll.save(update_fields=["status", "completion_date"])
        return Response(PollSerializer(poll, context={"request": request}).data)

    @action(
        detail=True, methods=["post"], serializer_class=PollSubmissionCreateSerializer
    )
    def submit_answers(self, request, pk=None):
        """Принимает ответы пользователя на опрос."""
        poll = get_object_or_404(Poll, pk=pk)

        can_submit = True
        error_message = ""

        if poll.status != Poll.StatusChoices.PUBLISHED:
            can_submit = False
            error_message = "Данный опрос не опубликован."
        elif poll.pub_date and poll.pub_date > timezone.now():
            can_submit = False
            error_message = "Данный опрос еще не доступен для прохождения."
        elif poll.completion_date and poll.completion_date < timezone.now():
            can_submit = False
            error_message = "Срок прохождения данного опроса истек."

        if can_submit and not poll.is_anonymous and request.user.is_authenticated:
            if PollSubmission.objects.filter(poll=poll, user=request.user).exists():
                can_submit = False
                error_message = "Вы уже проходили этот опрос."

        if can_submit and not poll.is_public:
            if not (
                request.user.is_authenticated
                and hasattr(request.user, "organization")
                and request.user.organization
                and poll.organization.filter(id=request.user.organization.id).exists()
            ):
                can_submit = False
                error_message = "Данный опрос недоступен для вашей организации."

        if not can_submit:
            return Response({"detail": error_message}, status=status.HTTP_403_FORBIDDEN)

        serializer = self.get_serializer(
            data=request.data, context={"request": request, "poll": poll}
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(
            {"message": "Ваши ответы успешно приняты."}, status=status.HTTP_201_CREATED
        )

    @action(detail=True, methods=["get"])
    def export(self, request, pk=None):
        """Возвращает статистику по опросу в формате XLSX."""
        poll = self.get_object()

        can_view_stats = False
        if request.user.is_staff or request.user.is_superuser:
            can_view_stats = True
        elif request.user.is_authenticated and (
            poll.author == request.user
            or poll.editors.filter(id=request.user.id).exists()
            or poll.stats_viewers.filter(id=request.user.id).exists()
        ):
            can_view_stats = True

        if not can_view_stats:
            return Response(
                {"detail": "У вас нет прав для просмотра статистики этого опроса."},
                status=status.HTTP_403_FORBIDDEN,
            )
        
        if poll.kind == Poll.KindChoices.FORM:
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="form_{poll.id}_answers.xlsx"'

            wb = Workbook()
            ws = wb.active
            ws.title = "Ответы на анкету"

            header_font = Font(bold=True)
            header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            data_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

            questions = poll.questions.order_by('order')
            headers = ["№"]
            question_headers = [q.text for q in questions]
            headers.extend(question_headers)

            for col_num, header_title in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header_title)
                cell.font = header_font
                cell.alignment = header_alignment
                cell.fill = header_fill
                cell.border = thin_border
            
            ws.row_dimensions[1].height = 30

            submissions = poll.submissions.prefetch_related(
                'answers__question',
                'answers__selected_choices'
            ).order_by('submitted_at')

            for row_idx, submission in enumerate(submissions, 2):
                col_idx = 1
                ws.cell(row=row_idx, column=col_idx, value=row_idx - 1).border = thin_border
                col_idx += 1

                submission_answers = {ans.question_id: ans for ans in submission.answers.all()}
                for question in questions:
                    answer = submission_answers.get(question.id)
                    answer_text = self._get_answer_text(answer, question)
                    cell = ws.cell(row=row_idx, column=col_idx, value=answer_text)
                    cell.alignment = data_alignment
                    cell.border = thin_border
                    col_idx += 1
            
            for col_idx, column_cells in enumerate(ws.columns, 1):
                max_length = 0
                column_letter = get_column_letter(col_idx)
                for cell in column_cells:
                    try:
                        if cell.value:
                            cell_text_len = max(len(line) for line in str(cell.value).split('\n'))
                            if cell_text_len > max_length:
                                max_length = cell_text_len
                    except:
                        pass
                adjusted_width = (max_length + 2) * 1.2
                ws.column_dimensions[column_letter].width = min(adjusted_width, 70)

            wb.save(response)
            return response

        else:

            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="poll_{poll.id}_statistics.xlsx"'

            wb = Workbook()
            ws = wb.active
            ws.title = "Статистика по опросу"

            header_font = Font(bold=True, size=12)
            question_header_font = Font(bold=True, italic=True, size=11)
            center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            left_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            thin_border = Border(left=Side(style='thin'), 
                                right=Side(style='thin'), 
                                top=Side(style='thin'), 
                                bottom=Side(style='thin'))
            
            row_num = 1
            ws.cell(row=row_num, column=1, value="ID Опроса").font = header_font
            ws.cell(row=row_num, column=2, value=poll.id)
            row_num += 1
            ws.cell(row=row_num, column=1, value="Название Опроса").font = header_font
            ws.cell(row=row_num, column=2, value=poll.name if poll.name else "N/A")
            row_num += 1
            total_submissions_count = poll.submissions.count()
            ws.cell(row=row_num, column=1, value="Всего прохождений").font = header_font
            ws.cell(row=row_num, column=2, value=total_submissions_count)
            row_num += 2

            questions_with_related = poll.questions.prefetch_related('choices', 'answers__selected_choices')

            for question in questions_with_related:
                ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
                q_header_cell = ws.cell(row=row_num, column=1, value=f"--- Вопрос ID: {question.id} ---")
                q_header_cell.font = question_header_font
                q_header_cell.alignment = center_alignment
                row_num += 1
                
                ws.cell(row=row_num, column=1, value="Текст вопроса:").font = Font(bold=True)
                ws.cell(row=row_num, column=2, value=question.text).alignment = left_alignment
                ws.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=4)
                row_num += 1
                
                ws.cell(row=row_num, column=1, value="Тип вопроса:").font = Font(bold=True)
                ws.cell(row=row_num, column=2, value=question.get_question_type_display())
                row_num += 1

                if question.question_type in [Question.QuestionType.SINGLE_CHOICE, Question.QuestionType.MULTIPLE_CHOICE]:
                    ws.cell(row=row_num, column=1, value="Вариант ответа").font = Font(bold=True)
                    ws.cell(row=row_num, column=2, value="Количество выборов").font = Font(bold=True)
                    ws.cell(row=row_num, column=3, value="Процент").font = Font(bold=True)
                    row_num += 1
                    
                    annotated_choices = question.choices.annotate(
                        num_answers=Count('chosen_in_answers', filter=Q(chosen_in_answers__submission__poll=poll))
                    )
                    for choice in annotated_choices:
                        count = choice.num_answers
                        percentage = (count / total_submissions_count * 100) if total_submissions_count > 0 else 0
                        ws.cell(row=row_num, column=1, value=choice.choice_text).alignment = left_alignment
                        ws.cell(row=row_num, column=2, value=count).alignment = center_alignment
                        ws.cell(row=row_num, column=3, value=f"{percentage:.2f}%").alignment = center_alignment
                        ws.cell(row=row_num, column=3).number_format = '0.00"%"'
                        row_num += 1
                    
                    if question.allow_custom_answer:
                        custom_answers = Answer.objects.filter(
                            question=question, submission__poll=poll
                        ).exclude(custom_choice_text__exact='').exclude(custom_choice_text__isnull=True)
                        
                        custom_answers_count = custom_answers.count()
                        custom_percentage = (custom_answers_count / total_submissions_count * 100) if total_submissions_count > 0 else 0
                        
                        ws.cell(row=row_num, column=1, value="Другое (свой вариант)").font = Font(bold=True)
                        ws.cell(row=row_num, column=2, value=custom_answers_count).alignment = center_alignment
                        ws.cell(row=row_num, column=3, value=f"{custom_percentage:.2f}%").alignment = center_alignment
                        ws.cell(row=row_num, column=3).number_format = '0.00"%"'
                        row_num += 1
                        
                        if custom_answers.exists():
                            ws.cell(row=row_num, column=1, value="Тексты своих вариантов ('Другое'):").font = Font(italic=True)
                            row_num += 1
                            for c_ans_text in custom_answers.values_list('custom_choice_text', flat=True):
                                ws.cell(row=row_num, column=1, value=c_ans_text).alignment = left_alignment
                                ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
                                row_num += 1
                
                elif question.question_type in [
                    Question.QuestionType.FREE_TEXT, 
                    Question.QuestionType.DATE, 
                    Question.QuestionType.TELEPHONE, 
                    Question.QuestionType.MAIL
                ]:
                    text_answers_qs = Answer.objects.filter(
                        question=question, 
                        submission__poll=poll
                    ).exclude(free_text_answer__exact='').exclude(free_text_answer__isnull=True)
                    
                    ws.cell(row=row_num, column=1, value="Текстовые ответы:").font = Font(bold=True)
                    ws.cell(row=row_num, column=2, value=text_answers_qs.count()).alignment = center_alignment
                    row_num += 1
                    
                    if text_answers_qs.exists():
                        for ans_text in text_answers_qs.values_list('free_text_answer', flat=True):
                            ws.cell(row=row_num, column=1, value=ans_text).alignment = left_alignment
                            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
                            row_num += 1
                    else:
                        ws.cell(row=row_num, column=1, value="Нет текстовых ответов")
                        row_num += 1
                
                row_num += 1

            for col_idx in range(1, 5):
                column_letter = get_column_letter(col_idx)
                max_length = 0
                for row_idx in range(1, row_num):
                    cell_value = ws[f"{column_letter}{row_idx}"].value
                    if cell_value:
                        if isinstance(cell_value, (int, float)):
                            cell_len = len(str(cell_value))
                        else:
                            cell_len = len(str(cell_value))
                        
                        lines = str(cell_value).split('\n')
                        max_line_len = max(len(line) for line in lines) if lines else 0
                        
                        if max_line_len > max_length:
                            max_length = max_line_len
                adjusted_width = (max_length + 2) * 1.2
                ws.column_dimensions[column_letter].width = min(adjusted_width, 70)

            wb.save(response)
            return response
    
    @action(detail=True, methods=['get'])
    def statistics(self, request, pk=None):
        """
        Возвращает агрегированную статистику по ответам на каждый вопрос опроса.
        """
        poll = self.get_object()

        can_view_stats = False
        if request.user.is_staff or request.user.is_superuser:
            can_view_stats = True
        elif request.user.is_authenticated and (
            poll.author == request.user
            or poll.editors.filter(id=request.user.id).exists()
            or poll.stats_viewers.filter(id=request.user.id).exists()
        ):
            can_view_stats = True
        
        if not can_view_stats:
            return Response(
                {"detail": "У вас нет прав для просмотра статистики этого опроса."},
                status=status.HTTP_403_FORBIDDEN,
            )

        total_submissions_count = poll.submissions.count()
        
        response_data = {
            "poll_id": poll.id,
            "poll_name": poll.name if poll.name else "N/A",
            "total_submissions": total_submissions_count,
            "question_statistics": []
        }

        questions_with_related = poll.questions.prefetch_related('choices', 'answers__selected_choices')

        for question in questions_with_related:
            q_stat = {
                "question_id": question.id,
                "text": question.text,
                "question_type": question.question_type,
                "question_type_display": question.get_question_type_display()
            }

            if question.question_type in [Question.QuestionType.SINGLE_CHOICE, Question.QuestionType.MULTIPLE_CHOICE]:
                q_stat["choices_stats"] = []
                
                annotated_choices = question.choices.annotate(
                    num_answers=Count('chosen_in_answers', filter=Q(chosen_in_answers__submission__poll=poll))
                )

                for choice in annotated_choices:
                    count = choice.num_answers
                    percentage = (count / total_submissions_count * 100) if total_submissions_count > 0 else 0
                    q_stat["choices_stats"].append({
                        "choice_id": choice.id,
                        "choice_text": choice.choice_text,
                        "count": count,
                        "percentage": round(percentage, 2)
                    })

                if question.allow_custom_answer:
                    custom_answers_qs = Answer.objects.filter(
                        question=question, submission__poll=poll
                    ).exclude(custom_choice_text__exact='').exclude(custom_choice_text__isnull=True)
                    
                    custom_answers_count = custom_answers_qs.count()
                    custom_percentage = (custom_answers_count / total_submissions_count * 100) if total_submissions_count > 0 else 0
                    
                    q_stat["custom_answers_stats"] = {
                        "label": "Другое (свой вариант)",
                        "count": custom_answers_count,
                        "percentage": round(custom_percentage, 2),
                        "sample_texts": list(custom_answers_qs.values_list('custom_choice_text', flat=True)[:5])
                    }

            elif question.question_type in [
                Question.QuestionType.FREE_TEXT, 
                Question.QuestionType.DATE, 
                Question.QuestionType.TELEPHONE, 
                Question.QuestionType.MAIL
            ]:
                text_answers_qs = Answer.objects.filter(
                    question=question, 
                    submission__poll=poll
                ).exclude(free_text_answer__exact='').exclude(free_text_answer__isnull=True)
                
                q_stat["free_text_answers_count"] = text_answers_qs.count()
                q_stat["sample_free_text_answers"] = list(text_answers_qs.values_list('free_text_answer', flat=True)[:5])
            
            response_data["question_statistics"].append(q_stat)
        
        return Response(response_data)


    @action(detail=True, methods=['get'], url_path='answers')
    def user_answers_list(self, request, pk=None):
        """
        Возвращает список ответов пользователей для данного опроса.
        Доступно только для неанонимных опросов.
        """
        poll = self.get_object()

        can_view_details = False
        if request.user.is_staff or request.user.is_superuser:
            can_view_details = True
        elif request.user.is_authenticated and (
            poll.author == request.user
            or poll.editors.filter(id=request.user.id).exists()
            or poll.stats_viewers.filter(id=request.user.id).exists()
        ):
            can_view_details = True
        
        if not can_view_details:
            return Response(
                {"detail": "У вас нет прав для просмотра детальных ответов этого опроса."},
                status=status.HTTP_403_FORBIDDEN,
            )

        if poll.is_anonymous:
            return Response(
                {"detail": "Просмотр ответов по пользователям недоступен для анонимных опросов."},
                status=status.HTTP_400_BAD_REQUEST
            )

        submissions = poll.submissions.filter(user__isnull=False).select_related('user').prefetch_related(
            'answers__question', 
            'answers__selected_choices'
        ).order_by('submitted_at')


        serializer = PollSubmissionWithAnswersSerializer(submissions, many=True, context={'request': request})
        return Response({
            'poll_id': poll.id,
            'poll_name': poll.name,
            'results': serializer.data
        })
    

    @action(detail=True, methods=['get'], url_path='answers/(?P<user_pk>[^/.]+)') # (?P<user_pk>[^/.]+) - для UUID или int
    def retrieve_user_answers(self, request, pk=None, user_pk=None):
        """
        Возвращает ответы конкретного пользователя на данный опрос.
        pk - ID опроса
        user_pk - ID пользователя
        """
        poll = self.get_object()

        if poll.is_anonymous:
            return Response(
                {"detail": "Просмотр ответов по пользователям недоступен для анонимных опросов."},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            target_user = get_object_or_404(Employee, pk=user_pk)
        except (ValueError, Employee.DoesNotExist):
             return Response(
                {"detail": "Пользователь с указанным ID не найден."},
                status=status.HTTP_404_NOT_FOUND
            )

        current_user = request.user
        can_view_target_user_answers = False

        if current_user.is_staff or current_user.is_superuser:
            can_view_target_user_answers = True
        elif current_user.is_authenticated:
            if (poll.author == current_user or
                poll.editors.filter(id=current_user.id).exists() or
                poll.stats_viewers.filter(id=current_user.id).exists()):
                can_view_target_user_answers = True
            elif current_user == target_user:
                can_view_target_user_answers = True
        
        if not can_view_target_user_answers:
            return Response(
                {"detail": "У вас нет прав для просмотра ответов этого пользователя на данный опрос."},
                status=status.HTTP_403_FORBIDDEN,
            )

        try:
            submission = PollSubmission.objects.select_related('user').prefetch_related(
                'answers__question', 
                'answers__selected_choices'
            ).get(poll=poll, user=target_user)
        except PollSubmission.DoesNotExist:
            return Response(
                {"detail": "Указанный пользователь не проходил данный опрос, или ответы не найдены."},
                status=status.HTTP_404_NOT_FOUND
            )

        serializer = PollSubmissionWithAnswersSerializer(submission, context={'request': request})
        
        response_data = {
            'poll_id': poll.id,
            'poll_name': poll.name,
            'submission_details': serializer.data
        }
        return Response(response_data)
    
    


class NewsViewSet(viewsets.ModelViewSet):
    """Вьюсет для новостей"""

    filter_backends = (DjangoFilterBackend,)

    filterset_fields = ("organization__id",)

    permission_classes = (
        IsAuthenticated,
        IsAdminUserOrReadOnly,
    )
    serializer_class = NewsSerializer

    def get_queryset(self):
        return News.objects.filter(
            is_published=True, pub_date__lte=timezone.now()
        ).order_by("-pub_date")


class ColleagueProfileViewset(UserViewSet):
    """Вьюсет для профиля"""

    permission_classes = (
        IsAuthenticated,
        IsAdminUserOrReadOnly,
    )
    queryset = Employee.objects.all()

    filter_backends = (DjangoFilterBackend, filters.SearchFilter)
    filterset_fields = (
        "structural_division__name",
        "structural_division__id",
        "chief__id",
        "structural_division__organization__id",
    )
    filterset_fields = {
        "structural_division__name": ["exact", "icontains"],
        "structural_division__id": ["exact", "isnull"],
        "chief__id": ["exact", "isnull"],
        "structural_division__organization__id": ["exact", "isnull"],
    }
    search_fields = ('email', 'surname', 'name', 'patronym', 'birth_date', 'email', 'telephone_number', 'inner_telephone_number', 'office', 'job_title', 'class_rank', 'status')

    def get_queryset(self):
        queryset = super().get_queryset()
        sort_by = self.request.query_params.get("sort_by")

        if sort_by == "name":
            queryset = queryset.annotate(
                full_name=Concat(
                    "surname",
                    Value(" "),
                    "name",
                    Value(" "),
                    "patronym",
                    Value(" "),
                    "email",
                    output_field=CharField(),
                )
            ).order_by("full_name")
        elif sort_by:
            valid_fields = [field.name for field in Employee._meta.fields]
            if sort_by in valid_fields:
                queryset = queryset.order_by(sort_by)
        return queryset

    def get_serializer_class(self):
        if self.action == "me":
            return MyProfileSerializer
        elif (
            self.action not in ("list", "retrieve", "create")
            and self.request.user.is_staff
            and self.kwargs.get("username")
            and self.kwargs.get("username") != self.request.user.username
        ):
            return ProfileInOrganizationSerializer

        return super().get_serializer_class()

    @staticmethod
    def validate_rating(serializer_class, request, id):
        user = request.user
        employee = get_object_or_404(Employee, id=id)
        request.data["user"] = user.id
        request.data["employee"] = employee.id

        serializer = serializer_class(
            data=request.data,
        )
        serializer.is_valid(raise_exception=True)

        return serializer

    @transaction.atomic
    @action(
        detail=True,
        methods=[
            "post",
        ],
        http_method_names=["post", "put", "delete"],
        permission_classes=(IsAuthenticated,),
        serializer_class=RatingPOSTSerializer,
    )
    def rate(self, request, id):
        serializer = self.validate_rating(RatingPOSTSerializer, request, id)
        employee = serializer.validated_data.get("employee")
        serializer.save()

        return Response(
            {
                "message": "Вы успешно оценили сотрудника.",
                "average_rating": employee.average_rating,
                "num_rates": employee.rated.count(),
            },
            status=status.HTTP_200_OK,
        )

    @transaction.atomic
    @rate.mapping.delete
    def unrate(self, request, id):
        serializer = self.validate_rating(RatingDELETESerializer, request, id)

        employee = serializer.validated_data.get("employee")

        Rating.objects.get(user=request.user, employee=employee).delete()

        return Response(
            {
                "message": "Вы успешно удалили свою оценку.",
                "average_rating": employee.average_rating,
                "num_rates": employee.rated.count(),
            },
            status=status.HTTP_204_NO_CONTENT,
        )

    @transaction.atomic
    @rate.mapping.put
    def change_or_rate(self, request, id):
        serializer = self.validate_rating(RatingPUTSerializer, request, id)
        employee = serializer.validated_data.get("employee")
        serializer.save()

        return Response(
            {
                "message": "Вы успешно оценили сотрудника.",
                "average_rating": employee.average_rating,
                "num_rates": employee.rated.count(),
            },
            status=status.HTTP_200_OK,
        )


class OrgStructureViewset(ListModelMixin, RetrieveModelMixin, viewsets.GenericViewSet):
    """Вьюсет для орг. структуры"""

    serializer_class = OrgStructureSerializer
    permission_classes = (
        IsAuthenticated,
        IsAdminUserOrReadOnly,
    )
    queryset = Employee.objects.all()


class OrganizationViewSet(viewsets.ModelViewSet):
    """Вьюсет для организаций для страницы Орг. структуры"""

    serializer_class = OrganizationSerializer
    permission_classes = (
        IsAuthenticated,
        IsAdminUserOrReadOnly,
    )
    queryset = Organization.objects.all()
    filter_backends = (DjangoFilterBackend, filters.SearchFilter)
    filterset_fields = (
        "name",
        "structural_subdivisions",
        "structural_subdivisions__positions__job_title",
        "structural_subdivisions__positions__class_rank",
        "structural_subdivisions__positions__status",
    )
    search_fields = (
        "structural_subdivisions__positions__name",
        "structural_subdivisions__positions__surname",
        "structural_subdivisions__positions__patronym",
    )


class HierarchyViewSet(ListModelMixin, RetrieveModelMixin, viewsets.GenericViewSet):
    """Вьюсет для иерархии."""

    filter_backends = (DjangoFilterBackend,)

    filterset_fields = ("id",)

    serializer_class = HierarchySerializer
    permission_classes = (IsAuthenticated,)
    queryset = Organization.objects.all()


class AgreeWithDataProcessingView(APIView):
    permission_classes = (IsAuthenticated,)

    def post(self, request):
        user = request.user
        user.agreed_with_data_processing = True
        user.save(update_fields=["agreed_with_data_processing"])
        return Response(
            {"message": "Согласие на обрабокту персональных данных отправлено."},
            status=status.HTTP_200_OK,
        )


class ValidateNextCloudView(APIView):
    """
    На эндпоинт приходит запрос с заголовком Authorization: Basic base64("{email}:{token}")
    Нужно вернуть 200 или 401, проверив токен Django (djoser).
    И добавить заголовок WWW-Authenticate: Basic realm="Nextcloud"
    """

    permission_classes = ()  # Allow any by default

    def get(self, request):
        auth_header = request.META.get("HTTP_AUTHORIZATION", "")
        unauthorized = Response(
            status=status.HTTP_401_UNAUTHORIZED,
            headers={"WWW-Authenticate": 'Basic realm="Nextcloud"'},
        )

        if not auth_header.startswith("Basic "):
            return unauthorized

        try:
            encoded = auth_header.split(" ", 1)[1]
            decoded = base64.b64decode(encoded).decode("utf-8")
            email, token_key = decoded.split(":", 1)
        except Exception:
            return unauthorized

        try:
            token = Token.objects.select_related("user").get(key=token_key)
        except Token.DoesNotExist:
            return unauthorized

        if token.user.email != email:
            return unauthorized

        return Response(
            {"message": "Token is valid."},
            status=status.HTTP_200_OK,
            headers={"WWW-Authenticate": 'Basic realm="Nextcloud"'},
        )


class CompetenceListView(generics.ListAPIView):
    """
    Provides a read-only list of all available Competences (id and name).

    Supports searching by competence name using the 'search' query parameter.
    Example: /api/v1/competences/?search=Python
    """

    serializer_class = CompetenceSerializer
    permission_classes = (IsAuthenticated,)

    filter_backends = (
        DjangoFilterBackend,
        filters.SearchFilter,
    )
    filterset_class = CompetenceFilter
    search_fields = ["name"]

    def get_queryset(self):
        """
        Annotate the queryset with the count of related characteristics.
        """
        queryset = Competence.objects.annotate(
            characteristic_count=Count("characteristic")
        ).order_by("name")

        return queryset.order_by("name")


import logging

logger = logging.getLogger(__name__)


class CustomTokenCreateView(TokenCreateView):
    def post(self, request, *args, **kwargs):
        response = super().post(request, *args, **kwargs)

        if response.status_code == status.HTTP_200_OK:
            try:
                auth_token_key = response.data.get("auth_token")

                if not auth_token_key:
                    return response

                try:
                    token_obj = Token.objects.select_related("user").get(
                        key=auth_token_key
                    )
                except Token.DoesNotExist:
                    return response

                user_email = token_obj.user.email

                if not user_email:
                    return response

                credentials_to_encode = f"{user_email}:{auth_token_key}"
                credentials_as_bytes = credentials_to_encode.encode("utf-8")

                encoded_credentials_bytes = base64.b64encode(credentials_as_bytes)

                encoded_credentials_str = encoded_credentials_bytes.decode("utf-8")

                response.set_cookie(
                    key="nextcloud_authorization",
                    value=encoded_credentials_str,
                    secure=False,
                    httponly=True,
                    path="/",
                    samesite="Lax",
                )

            except Exception as e:
                pass

        return response


class PollGroupListView(generics.ListAPIView):
    """
    View for PollGroup
    """

    serializer_class = PollGroupSerializer
    permission_classes = (IsAuthenticated,)

    filter_backends = (
        DjangoFilterBackend,
        filters.SearchFilter,
    )
    filterset_fields = {
        "name": ["exact", "icontains"],
    }
    search_fields = ["name"]

    queryset = PollGroup.objects.all()
