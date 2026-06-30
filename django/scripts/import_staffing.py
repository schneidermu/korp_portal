import argparse
import logging
import re
import shlex
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from django.conf import settings
from django.db import transaction
from openpyxl import load_workbook

from employees.models import Employee, Organization, StructuralSubdivision
from homepage.constants import CHARFIELD_LENGTH

logger = logging.getLogger(__name__)


DEFAULT_FILES = (
    "Центральный аппарат Центррегионводхоз.xlsx",
    "Штатная расстановка АКВА на 01.04.26.xlsx",
    "Штатная расстановка РОСНИИВХ на 01.02.2026 без ЗП.xlsx",
    "ФГБУ РосНИИВХ.xls",
    "штатная расстановка на 27.05.2026.docx",
    "Амурское БВУ.doc",
    "Верхне-Волжское БВУ.docx",
    "Двинско-Печорское.md",
    "Донское.md",
    "Енисейское.xls",
    "Западно-Каспийское.doc",
    "Камское.md",
    "Кубанское.doc",
    "Ленское.md",
    "Московско-Окское.doc",
    "Невско-Ладожское.docx",
    "Нижне-Волжское.xlsx",
    "Нижне-Обского БВУ.doc",
)

ORG_BY_FILE = {
    "Центральный аппарат Центррегионводхоз.xlsx": "Центррегионводхоз",
    "Штатная расстановка АКВА на 01.04.26.xlsx": "Акваинфотека",
    "Штатная расстановка РОСНИИВХ на 01.02.2026 без ЗП.xlsx": "РосНИИВХ",
    "ФГБУ РосНИИВХ.xls": "РосНИИВХ",
    "штатная расстановка на 27.05.2026.docx": "ЦА ФАВР",
    "Амурское БВУ.doc": "Амурское БВУ",
    "Верхне-Волжское БВУ.docx": "Верхне-Волжское БВУ",
    "Двинско-Печорское.md": "Двинско-Печорское БВУ",
    "Донское.md": "Донское БВУ",
    "Енисейское.xls": "Енисейское БВУ",
    "Западно-Каспийское.doc": "Западно-Каспийское БВУ",
    "Камское.md": "Камское БВУ",
    "Кубанское.doc": "Кубанское БВУ",
    "Ленское.md": "Ленское БВУ",
    "Московско-Окское.doc": "Московско-Окское БВУ",
    "Невско-Ладожское.docx": "Невско-Ладожское БВУ",
    "Нижне-Волжское.xlsx": "Нижне-Волжское БВУ",
    "Нижне-Обского БВУ.doc": "Нижне-Обское БВУ",
}

VACANCY_RE = re.compile(r"\bвакан[а-яёА-ЯЁ]+\b", re.IGNORECASE)
FULL_FIO_RE = re.compile(
    r"\b([А-ЯЁ][а-яё-]+)\s+([А-ЯЁ][а-яё-]+)\s+([А-ЯЁ][а-яё-]+)\b",
)
INITIALS_RE = re.compile(r"\b([А-ЯЁ][а-яё-]+)\s+([А-ЯЁ])\.?\s*([А-ЯЁ])\.?\b")
STAKE_RE = re.compile(r"\(?\b0[,.]5\b\)?")
NUM_RE = re.compile(r"^\d+(?:[,.]\d+)?$")
POSITION_HINTS = (
    "руковод",
    "начальник",
    "замест",
    "специалист",
    "консультант",
    "эксперт",
    "водитель",
    "сторож",
    "уборщ",
    "архивариус",
    "делопроизвод",
    "помощник",
    "главн",
    "ведущ",
    "старш",
    "инженер",
    "бухгалтер",
    "и.о.",
)
HEADER_MARKERS = (
    "№ п/п",
    "наименование должност",
    "ф.и.о",
    "фамилия",
    "кол-во штат",
    "количество штат",
    "ко-во штат",
    "факт",
)
ROSVODRESURSY_DOCX = "штатная расстановка на 27.05.2026.docx"
FGBU_ROSNIIVH_XLS = "ФГБУ РосНИИВХ.xls"


@dataclass(frozen=True)
class StaffRow:
    source: str
    sheet: str
    row_number: int
    organization: str
    subdivision: str
    position: str
    surname: str
    name: Optional[str]
    patronym: Optional[str]
    raw_name: str

    @property
    def display_name(self):
        parts = [self.surname, self.name, self.patronym]
        return " ".join(part for part in parts if part)


class ImportStats:
    def __init__(self):
        self.rows = 0
        self.updated = 0
        self.missing = 0
        self.ambiguous = 0
        self.skipped = 0
        self.created_orgs = 0
        self.created_subdivisions = 0


def run(*script_args):
    script_args = normalize_script_args(script_args)
    parser = argparse.ArgumentParser(
        description="Import staffing organizations and structural subdivisions.",
    )
    parser.add_argument(
        "--base-dir",
        default=str(settings.BASE_DIR.parent / "import"),
        help="Directory with staffing source files.",
    )
    parser.add_argument(
        "--file",
        dest="files",
        action="append",
        help="Specific staffing file path or file name. Can be passed multiple times.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse and match rows without writing to the database.",
    )
    parser.add_argument(
        "--allow-ambiguous-first",
        action="store_true",
        help="Update the first employee when several employees match the same FIO.",
    )
    args = parser.parse_args(script_args)

    base_dir = Path(args.base_dir)
    files = [Path(path) for path in args.files] if args.files else [Path(f) for f in DEFAULT_FILES]
    paths = [path if path.is_absolute() else base_dir / path for path in files]

    stats = ImportStats()
    with transaction.atomic():
        for path in paths:
            import_file(path, args, stats)

        if args.dry_run:
            transaction.set_rollback(True)

    logger.info(
        "Staffing import completed: rows=%s updated=%s missing=%s ambiguous=%s "
        "skipped=%s orgs_created=%s subdivisions_created=%s dry_run=%s",
        stats.rows,
        stats.updated,
        stats.missing,
        stats.ambiguous,
        stats.skipped,
        stats.created_orgs,
        stats.created_subdivisions,
        args.dry_run,
    )
    print(
        "Staffing import completed: "
        f"rows={stats.rows} updated={stats.updated} missing={stats.missing} "
        f"ambiguous={stats.ambiguous} skipped={stats.skipped} "
        f"orgs_created={stats.created_orgs} "
        f"subdivisions_created={stats.created_subdivisions} dry_run={args.dry_run}",
    )


def import_file(path, args, stats):
    if not path.exists():
        raise FileNotFoundError(f"Staffing file not found: {path}")
    if path.stat().st_size == 0:
        stats.skipped += 1
        logger.warning("Skipping empty staffing file: %s", path)
        return

    organization_name = ORG_BY_FILE.get(path.name, path.stem)
    logger.info("Importing staffing file: %s as organization '%s'", path, organization_name)

    try:
        rows = list(parse_workbook(path, organization_name))
    except Exception:
        stats.skipped += 1
        logger.exception("Failed to parse staffing file: %s", path)
        return

    stats.rows += len(rows)
    organization = get_or_create_organization(organization_name, args.dry_run, stats)

    subdivision_cache = {}
    for row in rows:
        subdivision = None
        if should_import_subdivision(row):
            subdivision = get_or_create_subdivision(
                organization,
                row.organization,
                row.subdivision,
                args.dry_run,
                stats,
                subdivision_cache,
            )
        update_employee(row, subdivision, args, stats)


def parse_workbook(path, organization_name):
    suffix = path.suffix.lower()
    if suffix == ".xls":
        if path.name == FGBU_ROSNIIVH_XLS:
            yield from parse_fgbu_rosniivh_xls(path, path.name, organization_name)
        else:
            yield from parse_bvu_xls(path, path.name, organization_name)
    elif suffix == ".docx":
        if path.name == ROSVODRESURSY_DOCX:
            yield from parse_rosvodresursy_docx(path, path.name, organization_name)
        else:
            yield from parse_bvu_docx(path, path.name, organization_name)
    elif suffix == ".doc":
        yield from parse_bvu_doc(path, path.name, organization_name)
    elif suffix == ".md":
        yield from parse_bvu_markdown(path, path.name, organization_name)
    elif suffix == ".pdf":
        yield from parse_bvu_pdf(path, path.name, organization_name)
    elif suffix == ".xlsx":
        workbook = load_workbook(path, data_only=True)
        if path.name.startswith("Центральный аппарат"):
            yield from parse_centerregionvodhoz(workbook.active, path.name, organization_name)
        elif "АКВА" in path.name:
            yield from parse_akva(workbook["Лист1"], path.name, organization_name)
        elif "РОСНИИВХ" in path.name.upper():
            yield from parse_rosniivh(workbook.active, path.name, organization_name)
        else:
            yield from parse_bvu_xlsx(workbook.active, path.name, organization_name)
    else:
        raise ValueError(f"Unknown staffing file format: {path}")


def parse_centerregionvodhoz(sheet, source, organization):
    current_subdivision = None
    for row_number in range(9, sheet.max_row + 1):
        unit = clean_text(sheet.cell(row_number, 2).value)
        position = clean_text(sheet.cell(row_number, 5).value)
        raw_name = clean_text(sheet.cell(row_number, 6).value)

        if unit and not is_total(unit):
            current_subdivision = unit

        if not position or not raw_name or is_vacancy(raw_name) or not current_subdivision:
            continue

        for person in extract_full_names(raw_name):
            yield StaffRow(source, sheet.title, row_number, organization, current_subdivision, position, *person, raw_name)


def parse_akva(sheet, source, organization):
    current_subdivision = None
    for row_number in range(4, sheet.max_row + 1):
        first = clean_text(sheet.cell(row_number, 1).value)
        position = clean_text(sheet.cell(row_number, 2).value)
        raw_name = clean_text(sheet.cell(row_number, 6).value)

        if first and not position and not raw_name and not is_total(first):
            current_subdivision = first
            continue

        if not position or is_total(position) or not raw_name or is_vacancy(raw_name):
            continue

        subdivision = current_subdivision or "Аппарат управления"
        for person in extract_full_names(raw_name):
            yield StaffRow(source, sheet.title, row_number, organization, subdivision, position, *person, raw_name)


def parse_rosniivh(sheet, source, organization):
    current_subdivision = "Аппарат управления"
    section_number = 1
    current_position = ""
    for row_number in range(12, sheet.max_row + 1):
        position = clean_text(sheet.cell(row_number, 1).value)
        raw_name = clean_text(sheet.cell(row_number, 3).value)

        if position and is_total(position):
            continue

        if not position and not raw_name:
            next_position = clean_text(sheet.cell(row_number + 1, 1).value)
            if next_position:
                section_number += 1
                current_subdivision = f"Подразделение {section_number}"
            continue

        if not raw_name or is_vacancy(raw_name):
            continue

        if position:
            current_position = position

        for person in extract_initial_names(raw_name):
            yield StaffRow(
                source,
                sheet.title,
                row_number,
                organization,
                current_subdivision,
                current_position,
                *person,
                raw_name,
            )


def parse_fgbu_rosniivh_xls(path, source, organization):
    import xlrd

    wb = xlrd.open_workbook(str(path))
    sh = wb.sheet_by_index(0)
    current_subdivision = None
    for row_number in range(4, sh.nrows):
        col0 = clean_text(sh.cell_value(row_number, 0))
        col1 = clean_text(sh.cell_value(row_number, 1))
        col2 = clean_text(sh.cell_value(row_number, 2))

        if not col1 and not col2:
            if col0 and not is_total(col0):
                current_subdivision = col0
            continue

        if not col1 or not col2 or is_vacancy(col2) or not current_subdivision:
            continue

        for person in extract_initial_names(col2):
            yield StaffRow(source, sh.name, row_number, organization, current_subdivision, col1, *person, col2)


def parse_rosvodresursy_docx(path, source, organization):
    from docx import Document

    doc = Document(str(path))
    table = doc.tables[0]
    current_subdivision = "Руководство"
    for row_number, row in enumerate(table.rows):
        if row_number < 2:
            continue

        cells = [clean_text(c.text) for c in row.cells]
        position = cells[1]
        fio = cells[3]

        # merged header row: position and FIO columns contain the same text
        if position and fio and position == fio:
            current_subdivision = position
            continue

        if not position or not fio or is_vacancy(fio):
            continue

        for person in extract_full_names(fio):
            yield StaffRow(source, "Лист1", row_number, organization, current_subdivision, position, *person, fio)


def parse_bvu_xlsx(sheet, source, organization):
    current_subdivision = "Аппарат управления"
    for row_number in range(1, sheet.max_row + 1):
        unit = clean_text(sheet.cell(row_number, 1).value)
        raw_name = clean_text(sheet.cell(row_number, 6).value)

        if unit and not raw_name and not is_total(unit):
            if is_subdivision_header(unit):
                current_subdivision = unit
            continue

        if not unit or is_total(unit) or not raw_name or is_vacancy(raw_name):
            continue

        for person in extract_people(raw_name):
            yield StaffRow(source, sheet.title, row_number, organization, current_subdivision, unit, *person, raw_name)


def parse_bvu_xls(path, source, organization):
    import xlrd

    sh = xlrd.open_workbook(str(path)).sheet_by_index(0)
    current_subdivision = "Аппарат управления"
    for row_number in range(6, sh.nrows):
        col0 = clean_text(sh.cell_value(row_number, 0))
        col1 = clean_text(sh.cell_value(row_number, 1))
        col3 = clean_text(sh.cell_value(row_number, 3))

        if col0 and not col1 and not col3 and not is_total(col0):
            current_subdivision = col0
            continue

        if not col0 or not col3 or is_vacancy(col3) or is_total(col0):
            continue

        seen = set()
        for person in extract_people(col3):
            key = (current_subdivision, col0, person)
            if key in seen:
                continue
            seen.add(key)
            yield StaffRow(source, sh.name, row_number, organization, current_subdivision, col0, *person, col3)


def parse_bvu_docx(path, source, organization):
    from docx import Document

    doc = Document(str(path))
    table = doc.tables[0]
    header = [clean_text(cell.text).lower() for cell in table.rows[0].cells]
    fio_column = next(
        (index for index, value in enumerate(header) if "фио" in value or "фамил" in value),
        1,
    )
    current_subdivision = "Аппарат управления"
    for row_number, row in enumerate(table.rows[1:], start=1):
        cells = [clean_text(cell.text) for cell in row.cells]
        if not any(cells):
            continue

        if len(set(cells)) == 1 and cells[0] and not is_total(cells[0]):
            if is_subdivision_header(cells[0]):
                current_subdivision = cells[0]
            continue

        position = cells[0]
        raw_name = cells[fio_column] if fio_column < len(cells) else ""
        if not position or not raw_name or is_vacancy(raw_name) or is_total(position):
            continue
        if is_subdivision_header(position):
            current_subdivision = position
            continue

        for person in extract_people(raw_name):
            yield StaffRow(source, "Лист1", row_number, organization, current_subdivision, position, *person, raw_name)


def parse_bvu_doc(path, source, organization):
    lines = extract_legacy_doc_lines(path)
    yield from parse_bvu_text_lines(lines, source, organization)


def parse_bvu_markdown(path, source, organization):
    lines = path.read_text(encoding="utf-8").splitlines()
    table_lines = [line for line in lines if line.strip().startswith("|")]
    if len(table_lines) < 2:
        raise ValueError(f"No markdown table found in {path}")

    header_cells = parse_markdown_row(table_lines[0])
    if not header_cells:
        raise ValueError(f"Invalid markdown table header in {path}")

    header = [cell.lower() for cell in header_cells]
    fio_column = next(
        (index for index, value in enumerate(header) if "фио" in value or "фамил" in value),
        max(len(header_cells) - 2, 1),
    )

    current_subdivision = "Аппарат управления"
    for row_number, line in enumerate(table_lines[2:], start=3):
        if is_markdown_separator(line):
            continue

        raw_cells = parse_markdown_row(line, keep_markdown=True)
        cells = parse_markdown_row(line)
        if not cells:
            continue

        position = cells[0]
        if not position:
            continue

        if is_markdown_subdivision_row(cells, raw_cells, fio_column):
            current_subdivision = normalize_subdivision_name(position)
            continue

        if is_total(position):
            continue

        raw_name = cells[fio_column] if fio_column < len(cells) else ""
        if not raw_name or is_vacancy(raw_name):
            continue

        for person in extract_people(raw_name):
            yield StaffRow(source, path.stem, row_number, organization, current_subdivision, position, *person, raw_name)


def parse_bvu_pdf(path, source, organization):
    lines = extract_pdf_lines(path)
    yield from parse_bvu_text_lines(lines, source, organization)


def parse_bvu_text_lines(lines, source, organization):
    current_subdivision = "Аппарат управления"
    start = find_table_start(lines)
    index = start
    while index < len(lines):
        line = lines[index]
        if is_header_line(line) or line.lower().startswith(("уважаем", "приложение:")):
            index += 1
            continue
        if is_total(line) or re.match(r"^Штат\s*[–-]", line):
            index += 1
            continue
        if is_subdivision_header(line):
            current_subdivision = clean_text(line).rstrip(":")
            index += 1
            continue

        if NUM_RE.match(line) and index + 2 < len(lines) and looks_like_position(lines[index + 1]):
            position = lines[index + 1]
            raw_name = lines[index + 2]
            if not is_vacancy(raw_name):
                for person in extract_people(raw_name):
                    yield make_staff_row(source, index, organization, current_subdivision, position, person, raw_name)
            index += 3
            continue

        numbered_position = re.match(r"^\d+\.\s*(.+)$", line)
        if numbered_position and index + 1 < len(lines) and is_person_line(lines[index + 1]):
            position = numbered_position.group(1).strip()
            raw_name = lines[index + 1]
            for person in extract_people(raw_name):
                yield make_staff_row(source, index, organization, current_subdivision, position, person, raw_name)
            index += 2
            continue

        if looks_like_position(line):
            position_parts = [line]
            next_index = index + 1
            while next_index < len(lines) and looks_like_position(lines[next_index]) and not NUM_RE.match(lines[next_index]):
                position_parts.append(lines[next_index])
                next_index += 1
            position = " ".join(clean_text(part) for part in position_parts)
            people, next_index = consume_people(lines, next_index)
            for person, raw_name in people:
                yield make_staff_row(source, index, organization, current_subdivision, position, person, raw_name)
            index = next_index
            continue

        index += 1


def make_staff_row(source, row_number, organization, subdivision, position, person, raw_name):
    return StaffRow(source, "Лист1", row_number, organization, subdivision, position, *person, raw_name)


def consume_people(lines, start_index):
    people = []
    index = start_index
    while index < len(lines):
        line = lines[index]
        if is_subdivision_header(line) or is_total(line) or looks_like_position(line) or re.match(r"^\d+\.", line):
            break
        if NUM_RE.match(clean_text(line)):
            index += 1
            if index < len(lines) and is_person_line(lines[index]):
                raw_name = lines[index]
                people.extend((person, raw_name) for person in extract_people(raw_name))
                index += 1
            continue
        if is_vacancy(line):
            index += 1
            break
        if is_person_line(line):
            raw_name = line
            people.extend((person, raw_name) for person in extract_people(raw_name))
            index += 1
            continue
        break
    return people, index


def find_table_start(lines):
    for index, line in enumerate(lines):
        if clean_text(line) == "Наименование должностей":
            return index
    for index, line in enumerate(lines):
        lower = line.lower()
        if "штатная расстановка" in lower or "штатное замещение" in lower:
            return index
    return 0


def extract_legacy_doc_lines(path):
    for command in (["antiword", str(path)], ["textutil", "-convert", "txt", "-stdout", str(path)]):
        try:
            text = subprocess.check_output(command, text=True, stderr=subprocess.DEVNULL)
        except (FileNotFoundError, subprocess.CalledProcessError):
            continue
        return [line.strip() for line in text.splitlines() if line.strip()]
    raise RuntimeError(
        f"Cannot read legacy .doc file {path.name}: install antiword or use macOS textutil",
    )


def parse_markdown_row(line, keep_markdown=False):
    stripped = line.strip()
    if not stripped.startswith("|"):
        return None
    cells = stripped.strip("|").split("|")
    if keep_markdown:
        return [cell.strip() for cell in cells]
    return [clean_markdown_cell(cell) for cell in cells]


def clean_markdown_cell(cell):
    value = clean_text(cell)
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    return value.strip()


def is_markdown_separator(line):
    stripped = line.strip()
    if not stripped.startswith("|"):
        return False
    return set(stripped.replace("|", "").replace("-", "").replace(":", "").strip()) == set()


def is_markdown_subdivision_row(cells, raw_cells, fio_column):
    position = cells[0]
    if not position or is_total(position):
        return False
    raw_name = cells[fio_column] if fio_column < len(cells) else ""
    if clean_text(raw_name):
        return False
    if "**" in raw_cells[0]:
        return True
    return is_subdivision_header(position)


def normalize_subdivision_name(name):
    normalized = clean_text(name)
    address_start = normalized.find(" (")
    if address_start > 0:
        normalized = normalized[:address_start].strip()
    return truncate_subdivision_name(normalized)


def truncate_subdivision_name(name):
    return truncate_field(clean_text(name), "subdivision name")


def truncate_field(value, field_name, max_length=CHARFIELD_LENGTH):
    normalized = clean_text(value)
    if len(normalized) <= max_length:
        return normalized
    logger.warning(
        "Truncating %s to %s chars: %s",
        field_name,
        max_length,
        normalized,
    )
    return normalized[:max_length].rstrip()


def extract_pdf_lines(path):
    import shutil

    import pdfplumber

    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    text = "\n".join(text_parts)
    if clean_text(text):
        return [line.strip() for line in text.splitlines() if line.strip()]

    if not shutil.which("tesseract"):
        raise RuntimeError(
            f"Scanned PDF {path.name} requires tesseract OCR (tesseract-ocr + tesseract-ocr-rus)",
        )

    from pdf2image import convert_from_path
    import pytesseract

    ocr_parts = []
    for image in convert_from_path(str(path), dpi=200):
        ocr_parts.append(pytesseract.image_to_string(image, lang="rus"))
    return [line.strip() for line in "\n".join(ocr_parts).splitlines() if line.strip()]


def extract_people(value):
    people = list(extract_full_names(value))
    if people:
        return people
    return list(extract_initial_names(value))


def looks_like_position(text):
    normalized = clean_text(text).lower()
    if not normalized or is_total(normalized) or NUM_RE.match(normalized):
        return False
    return any(hint in normalized for hint in POSITION_HINTS)


def is_subdivision_header(text):
    normalized = clean_text(text)
    if not normalized or is_total(normalized) or looks_like_position(normalized) or NUM_RE.match(normalized):
        return False
    lower = normalized.lower().rstrip(":")
    if lower in ("руководство",):
        return True
    if lower.startswith(("отдел ", "аппарат ", "территориаль", "персонал по")):
        return True
    return "отдел " in lower


def is_header_line(line):
    normalized = clean_text(line).lower()
    return any(marker in normalized for marker in HEADER_MARKERS) and not looks_like_position(line)


def is_person_line(line):
    if not line or is_vacancy(line) or is_total(line) or is_subdivision_header(line):
        return False
    if NUM_RE.match(clean_text(line)):
        return False
    return bool(extract_people(line))


def get_or_create_organization(name, dry_run, stats):
    if dry_run:
        organization = Organization.objects.filter(name=name).first()
        if organization is None:
            logger.info("DRY-RUN would create organization: %s", name)
            stats.created_orgs += 1
        return organization

    organization, created = Organization.objects.get_or_create(
        name=name,
        defaults={"address": ""},
    )
    if created:
        stats.created_orgs += 1
        logger.info("Created organization: %s", name)
    return organization


def get_or_create_subdivision(
    organization,
    organization_name,
    subdivision_name,
    dry_run,
    stats,
    cache,
):
    subdivision_name = truncate_subdivision_name(subdivision_name)
    key = (organization_name, subdivision_name)
    if key in cache:
        return cache[key]

    if dry_run:
        subdivision = None
        if organization is not None:
            subdivision = StructuralSubdivision.objects.filter(
                organization=organization,
                name=subdivision_name,
            ).first()
        if subdivision is None:
            logger.info(
                "DRY-RUN would create subdivision: %s / %s",
                organization_name,
                subdivision_name,
            )
            stats.created_subdivisions += 1
        cache[key] = subdivision
        return subdivision

    subdivision, created = StructuralSubdivision.objects.get_or_create(
        organization=organization,
        name=subdivision_name,
    )
    if created:
        stats.created_subdivisions += 1
        logger.info("Created subdivision: %s / %s", organization_name, subdivision_name)

    cache[key] = subdivision
    return subdivision


def update_employee(row, subdivision, args, stats):
    matches = find_employee(row)
    if not matches:
        stats.missing += 1
        logger.warning(
            "missing: %s | %s | %s | row=%s | raw=%s",
            row.display_name,
            row.organization,
            row.subdivision,
            row.row_number,
            row.raw_name,
        )
        return

    if len(matches) > 1 and not args.allow_ambiguous_first:
        stats.ambiguous += 1
        logger.warning(
            "ambiguous: %s | matches=%s | %s | %s | row=%s",
            row.display_name,
            len(matches),
            row.organization,
            row.subdivision,
            row.row_number,
        )
        return

    employee = matches[0]
    if subdivision is None and not employee_has_target_organization(employee, row.organization):
        stats.missing += 1
        logger.warning(
            "missing subdivision source: %s | %s | row=%s | current_unit=%s",
            row.display_name,
            row.organization,
            row.row_number,
            employee.structural_division,
        )
        return

    if args.dry_run:
        stats.updated += 1
        logger.info(
            "DRY-RUN would update: %s -> %s / %s, position=%s",
            employee,
            row.organization,
            subdivision or employee.structural_division,
            row.position,
        )
        return

    if subdivision is not None:
        employee.structural_division = subdivision
    employee.job_title = truncate_field(row.position, "job_title")
    update_fields = ["job_title"]
    if subdivision is not None:
        update_fields.append("structural_division")
    employee.save(update_fields=update_fields)
    stats.updated += 1
    logger.info(
        "Updated: %s -> %s / %s",
        employee,
        row.organization,
        subdivision or employee.structural_division,
    )


def should_import_subdivision(row):
    if row.organization == "РосНИИВХ" and row.subdivision.startswith("Подразделение "):
        return False
    return True


def employee_has_target_organization(employee, organization_name):
    structural_division = employee.structural_division
    if structural_division is None:
        return False
    return structural_division.organization.name == organization_name


def find_employee(row):
    if uses_full_fio(row):
        return list(
            Employee.objects.filter(
                surname__iexact=row.surname,
                name__iexact=row.name,
                patronym__iexact=row.patronym,
            )
            .select_related("structural_division__organization")
            .order_by("id"),
        )

    candidates = [
        employee
        for employee in Employee.objects.filter(surname__iexact=row.surname)
        .select_related("structural_division__organization")
        .order_by("id")
        if initials_match(employee, row.name, row.patronym)
    ]
    if not candidates:
        return []

    org_candidates = [
        employee
        for employee in candidates
        if employee_has_target_organization(employee, row.organization)
    ]
    if len(org_candidates) == 1:
        return org_candidates
    if len(org_candidates) > 1:
        return org_candidates

    if len(candidates) == 1:
        return candidates
    return candidates


def uses_full_fio(row):
    return (
        row.name
        and row.patronym
        and len(row.name) > 1
        and len(row.patronym) > 1
    )


def initials_match(employee, name_initial, patronym_initial):
    employee_name = employee.name or ""
    employee_patronym = employee.patronym or ""
    return (
        employee_name[:1].upper() == (name_initial or "").upper()
        and employee_patronym[:1].upper() == (patronym_initial or "").upper()
    )


def extract_full_names(value):
    value = strip_notes(value)
    for match in FULL_FIO_RE.finditer(value):
        yield match.group(1), match.group(2), match.group(3)


def extract_initial_names(value):
    value = strip_notes(value)
    value = STAKE_RE.sub(" ", value)
    for match in INITIALS_RE.finditer(value):
        yield match.group(1), match.group(2), match.group(3)


def strip_notes(value):
    return re.sub(r"\([^)]*\)", " ", value)


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def is_total(value):
    return clean_text(value).lower().startswith(("итого", "всего"))


def is_vacancy(value):
    return bool(VACANCY_RE.search(clean_text(value)))


def normalize_script_args(script_args):
    if len(script_args) == 1 and isinstance(script_args[0], str):
        return shlex.split(script_args[0])
    return list(script_args)
